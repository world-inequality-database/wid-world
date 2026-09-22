#!/usr/bin/env python3
"""Export Excel sheets separately or join their columns on Grade.

Requires Python 3.9+ and openpyxl: python3 -m pip install openpyxl
Examples (default workbook: grade_definitions.xlsx beside this script):
  python3 export_grade_definitions.py --mode batch
  python3 export_grade_definitions.py --list-sheets
  python3 export_grade_definitions.py --sheets 1,2 --output merged.csv
  python3 export_grade_definitions.py --sheets 3 --output aggregate.csv
  python3 export_grade_definitions.py --sheets all --mode separate --output exports

Selections accept comma-separated sheet numbers (1-based) or exact names.
Merge keeps every observed Grade, sorted, leaving missing matches blank.
Merged columns use Short desc.: or Long desc.: prefixes.
Separate exports keep source headers and use full descriptive filenames.
Header whitespace is stripped; definition text and literal N/A are preserved.
CSVs use UTF-8 with BOM for Excel. Existing output CSVs are replaced.
The workbook is never modified. Formulas must be converted to values first.
"""
import argparse
import csv
import re
from pathlib import Path
from openpyxl import load_workbook


def read_sheet(sheet):
    rows = list(sheet.iter_rows())
    used = [i for i in range(sheet.max_column)
            if any(row[i].value is not None for row in rows)]
    headers = [str(rows[0][i].value or '').strip() for i in used]
    if not headers or '' in headers or len(set(headers)) != len(headers):
        raise ValueError(f'{sheet.title}: missing or duplicate headers')
    if 'Grade' not in headers:
        raise ValueError(f'{sheet.title}: no Grade column')
    key = headers.index('Grade')
    columns = [h for h in headers if h != 'Grade']
    data = {}
    for row in rows[1:]:
        cells = [row[i] for i in used]
        if all(c.value is None for c in cells):
            continue
        for cell in cells:
            if cell.data_type in ('f', 'e'):
                raise ValueError(f'{sheet.title}!{cell.coordinate}: formula or Excel error; literal values required')
        values = [c.value for c in cells]
        grade = values[key]
        if isinstance(grade, bool) or not isinstance(grade, (int, float)) or grade not in range(6):
            raise ValueError(f'{sheet.title}, row {cells[0].row}: Grade must be an integer from 0 to 5')
        grade = int(grade)
        if grade in data:
            raise ValueError(f'{sheet.title}: duplicate Grade {grade}')
        data[grade] = [v for i, v in enumerate(values) if i != key]
    if not data:
        raise ValueError(f'{sheet.title}: no grade rows')
    return columns, data


def select_sheets(selection, names):
    if selection.strip().lower() == 'all':
        return names[:]
    chosen = []
    for token in selection.split(','):
        token = token.strip()
        if token in names:
            name = token
        elif token.isdigit() and 1 <= int(token) <= len(names):
            name = names[int(token) - 1]
        else:
            raise ValueError(f'Unknown sheet {token!r}; use --list-sheets')
        if name in chosen:
            raise ValueError(f'Sheet selected twice: {name}')
        chosen.append(name)
    return chosen


def export_headers(sheet_name, columns):
    """Use consistent labels even when Excel truncates a sheet name."""
    match = re.search(r"\b(short|long)\b", sheet_name, flags=re.IGNORECASE)
    prefix = f"{match.group(1).capitalize()} desc." if match else sheet_name
    return [f"{prefix}: {column}" for column in columns]


def separate_filename(sheet_name):
    """Expand the known definition sheet names without numeric prefixes."""
    match = re.match(r"^(Distributed|Aggregated?) variables\s*\((short|long)\b", sheet_name, re.IGNORECASE)
    if match:
        group = 'distributed' if match.group(1).lower() == 'distributed' else 'aggregate'
        return f'{group}_variables_{match.group(2).lower()}_descriptions.csv'
    return (re.sub(r'[^\w.-]+', '_', sheet_name).strip('._') or 'sheet').lower() + '.csv'


def merge_tables(tables):
    headers = ['Grade'] + [header for name, (columns, _) in tables.items()
                          for header in export_headers(name, columns)]
    if len(set(headers)) != len(headers):
        raise ValueError('Merged headers collide; rename conflicting source headers')
    grades = sorted(set().union(*(data for _, data in tables.values())))
    rows = [[grade] + [v for columns, data in tables.values()
                      for v in data.get(grade, [None] * len(columns))] for grade in grades]
    return headers, rows


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--workbook', type=Path, default=Path(__file__).resolve().with_name('grade_definitions.xlsx'))
    parser.add_argument('--list-sheets', action='store_true')
    parser.add_argument('--sheets', help='Comma-separated sheet numbers or names, or all')
    parser.add_argument('--mode', choices=('merge', 'separate', 'batch'), default='merge')
    parser.add_argument('--output', type=Path, help='CSV filename for merge; directory for separate/batch')
    args = parser.parse_args()
    try:
        workbook = load_workbook(args.workbook, read_only=True, data_only=False)
        try:
            sheetnames = workbook.sheetnames
            if args.list_sheets:
                for number, name in enumerate(sheetnames, 1):
                    print(f'{number}: {name}')
                return
            if args.mode == 'batch' and args.sheets:
                parser.error('--mode batch exports the nine standard tables; omit --sheets')
            if not args.sheets and args.mode != 'batch':
                parser.error('--sheets is required unless using --list-sheets')
            names = select_sheets('all' if args.mode == 'batch' else args.sheets, sheetnames)
            tables = {name: read_sheet(workbook[name]) for name in names}
        finally:
            workbook.close()
        outputs = []
        if args.mode == 'merge':
            headers, rows = merge_tables(tables)
            outputs.append((args.output or Path('grade_definitions.csv'), headers, rows))
        else:
            directory = args.output or (args.workbook.parent / 'csv_outputs' if args.mode == 'batch' else Path('grade_exports'))
            if args.mode == 'batch':
                by_filename = {separate_filename(name): name for name in tables}
                required = [f'{group}_variables_{length}_descriptions.csv'
                            for group in ('aggregate', 'distributed') for length in ('long', 'short')]
                if len(tables) != 4 or set(by_filename) != set(required):
                    raise ValueError('Batch export requires exactly the four aggregate/distributed short/long definition sheets')
            for name, (columns, data) in tables.items():
                outputs.append((directory / separate_filename(name), ['Grade'] + columns,
                                [[grade] + data[grade] for grade in sorted(data)]))
            if args.mode == 'batch':
                combinations = {
                    'aggregate_variables_long_short.csv': required[:2],
                    'distributed_variables_long_short.csv': required[2:],
                    'short_descriptions_all.csv': [required[1], required[3]],
                    'long_descriptions_all.csv': [required[0], required[2]],
                    'all_descriptions_merged.csv': required,
                }
                for filename, sources in combinations.items():
                    selected = {by_filename[source]: tables[by_filename[source]] for source in sources}
                    headers, rows = merge_tables(selected)
                    outputs.append((directory / filename, headers, rows))
        if len({str(path.resolve()).casefold() for path, _, _ in outputs}) != len(outputs):
            raise ValueError('Selected sheets produce duplicate CSV filenames; rename the conflicting sheets')
        for path, _, _ in outputs:
            if path.resolve() == args.workbook.resolve():
                raise ValueError('Output must not overwrite the input workbook')
            if path.suffix.lower() != '.csv':
                raise ValueError('Output filename must end in .csv')
        for path, headers, rows in outputs:
            path.parent.mkdir(parents=True, exist_ok=True)
            with path.open('w', encoding='utf-8-sig', newline='') as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                writer.writerows(rows)
            print(f'Created {path}')
    except (ValueError, OSError) as error:
        parser.exit(1, f'Error: {error}\n')


if __name__ == '__main__':
    main()
